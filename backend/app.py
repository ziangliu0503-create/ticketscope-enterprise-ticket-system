import csv
import os
import sqlite3
from datetime import datetime, timedelta, timezone
from functools import wraps
from io import BytesIO, StringIO
from pathlib import Path

from flask import Flask, g, jsonify, request, send_file, send_from_directory
from flask_cors import CORS
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.security import check_password_hash, generate_password_hash

from .db import get_db, init_app as init_db_app, init_db
from .seed import seed_database


STATUSES = ["NEW", "ASSIGNED", "IN_PROGRESS", "RESOLVED", "CLOSED"]
PRIORITIES = ["LOW", "MEDIUM", "HIGH", "URGENT"]
DEPARTMENTS = ["财务部", "供应链部", "市场部", "销售部", "人力资源部", "生产运营部"]
CATEGORIES = ["账号权限", "网络连接", "软件故障", "数据问题", "设备故障", "系统咨询"]
SLA_HOURS = {"LOW": 72, "MEDIUM": 48, "HIGH": 24, "URGENT": 8}
STATUS_LABELS = {"NEW": "待受理", "ASSIGNED": "已分配", "IN_PROGRESS": "处理中", "RESOLVED": "已解决", "CLOSED": "已关闭"}
PRIORITY_LABELS = {"LOW": "低", "MEDIUM": "中", "HIGH": "高", "URGENT": "紧急"}
ROLE_LABELS = {"requester": "普通员工", "agent": "技术人员", "admin": "管理员"}

ALLOWED_TRANSITIONS = {
    "NEW": {"ASSIGNED"},
    "ASSIGNED": {"IN_PROGRESS"},
    "IN_PROGRESS": {"RESOLVED"},
    "RESOLVED": {"CLOSED", "IN_PROGRESS"},
    "CLOSED": {"IN_PROGRESS"},
}


def utc_now():
    return datetime.now(timezone.utc).replace(microsecond=0)


def iso(value):
    return value.astimezone(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def parse_iso(value):
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def row_to_dict(row):
    return dict(row) if row is not None else None


def error(message, status=400, details=None):
    payload = {"error": message}
    if details:
        payload["details"] = details
    return jsonify(payload), status


def sla_status(ticket, now=None):
    now = now or utc_now()
    due_at = parse_iso(ticket["due_at"])
    if ticket["status"] in {"RESOLVED", "CLOSED"}:
        completed_at = parse_iso(ticket["resolved_at"]) if ticket.get("resolved_at") else now
        return "MET" if completed_at <= due_at else "BREACHED"
    if now > due_at:
        return "OVERDUE"
    warning_hours = SLA_HOURS[ticket["priority"]] * 0.25
    return "WARNING" if due_at - now <= timedelta(hours=warning_hours) else "ON_TRACK"


def enrich_ticket(row):
    ticket = row_to_dict(row)
    if not ticket:
        return None
    ticket["sla_status"] = sla_status(ticket)
    remaining = (parse_iso(ticket["due_at"]) - utc_now()).total_seconds() / 3600
    ticket["sla_remaining_hours"] = round(remaining, 1)
    return ticket


def user_scope(user, alias="t"):
    if user["role"] == "admin":
        return "", []
    if user["role"] == "requester":
        return f"{alias}.requester_id = ?", [user["id"]]
    return f"({alias}.assignee_id = ? OR ({alias}.status = 'NEW' AND {alias}.assignee_id IS NULL))", [user["id"]]


def fetch_ticket(ticket_id, user=None):
    conditions = ["t.id = ?"]
    parameters = [ticket_id]
    if user:
        condition, values = user_scope(user)
        if condition:
            conditions.append(condition)
            parameters.extend(values)
    return get_db().execute(
        f"""
        SELECT t.*, requester.name AS requester_name, requester.email AS requester_email,
               assignee.name AS assignee_name
        FROM tickets t
        JOIN users requester ON requester.id = t.requester_id
        LEFT JOIN users assignee ON assignee.id = t.assignee_id
        WHERE {' AND '.join(conditions)}
        """,
        parameters,
    ).fetchone()


def allowed_transitions_for(ticket, user):
    state = ticket["status"]
    if user["role"] == "admin":
        return sorted(ALLOWED_TRANSITIONS.get(state, set()), key=STATUSES.index)
    if user["role"] == "agent":
        if state == "NEW" and ticket["assignee_id"] is None:
            return ["ASSIGNED"]
        if ticket["assignee_id"] == user["id"] and state in {"ASSIGNED", "IN_PROGRESS"}:
            return sorted(ALLOWED_TRANSITIONS[state], key=STATUSES.index)
        return []
    if ticket["requester_id"] == user["id"] and state == "RESOLVED":
        return ["CLOSED", "IN_PROGRESS"]
    if ticket["requester_id"] == user["id"] and state == "CLOSED":
        return ["IN_PROGRESS"]
    return []


def create_app(test_config=None):
    frontend_dist = Path(__file__).parents[1] / "frontend" / "dist"
    app = Flask(__name__, static_folder=str(frontend_dist) if frontend_dist.exists() else None, static_url_path="")
    default_database = Path(__file__).with_name("data") / "tickets.db"
    app.config.from_mapping(
        DATABASE=os.environ.get("TICKET_DATABASE", str(default_database)),
        SECRET_KEY=os.environ.get("TICKET_SECRET_KEY", "ticket-scope-demo-secret-change-in-production"),
        AUTH_TOKEN_MAX_AGE=int(os.environ.get("AUTH_TOKEN_MAX_AGE", "86400")),
        JSON_AS_ASCII=False,
    )
    if test_config:
        app.config.update(test_config)
    CORS(app)
    init_db_app(app)

    with app.app_context():
        init_db()
        db = get_db()
        if os.environ.get("TICKET_SEED_DEMO", "false").lower() == "true" and db.execute("SELECT COUNT(*) FROM tickets").fetchone()[0] == 0:
            seed_database(app.config["DATABASE"], count=220)
        demo_hash = generate_password_hash("Demo123!")
        db.execute("UPDATE users SET password_hash = ? WHERE password_hash IS NULL OR password_hash = ''", (demo_hash,))
        db.commit()

    def serializer():
        return URLSafeTimedSerializer(app.config["SECRET_KEY"], salt="ticket-scope-auth")

    def auth_required(*roles):
        def decorator(view):
            @wraps(view)
            def wrapped(*args, **kwargs):
                header = request.headers.get("Authorization", "")
                if not header.startswith("Bearer "):
                    return error("请先登录", 401)
                try:
                    payload = serializer().loads(header[7:], max_age=app.config["AUTH_TOKEN_MAX_AGE"])
                except SignatureExpired:
                    return error("登录已过期，请重新登录", 401)
                except BadSignature:
                    return error("登录凭证无效", 401)
                user = get_db().execute(
                    "SELECT id, name, email, department, role FROM users WHERE id = ?", (payload.get("user_id"),)
                ).fetchone()
                if user is None:
                    return error("用户不存在", 401)
                g.current_user = row_to_dict(user)
                if roles and g.current_user["role"] not in roles:
                    return error("当前角色没有此操作权限", 403)
                return view(*args, **kwargs)
            return wrapped
        return decorator

    def build_ticket_filters(user):
        conditions = []
        parameters = []
        scope, scope_values = user_scope(user)
        if scope:
            conditions.append(scope)
            parameters.extend(scope_values)
        for field in ("status", "priority", "department", "category"):
            value = request.args.get(field)
            if value:
                conditions.append(f"t.{field} = ?")
                parameters.append(value)
        search = request.args.get("search", "").strip()
        if search:
            conditions.append("(t.ticket_no LIKE ? OR t.title LIKE ? OR t.description LIKE ?)")
            token = f"%{search}%"
            parameters.extend([token, token, token])
        sla = request.args.get("sla", "")
        now = utc_now()
        if sla == "OVERDUE":
            conditions.append("t.status NOT IN ('RESOLVED', 'CLOSED') AND t.due_at < ?")
            parameters.append(iso(now))
        elif sla == "WARNING":
            warning_parts = []
            for priority, hours in SLA_HOURS.items():
                warning_parts.append("(t.priority = ? AND t.due_at BETWEEN ? AND ?)")
                parameters.extend([priority, iso(now), iso(now + timedelta(hours=hours * 0.25))])
            conditions.append("t.status NOT IN ('RESOLVED', 'CLOSED') AND (" + " OR ".join(warning_parts) + ")")
        elif sla == "ESCALATED":
            conditions.append("t.escalation_level > 0")
        return conditions, parameters

    def ticket_query(user, include_pagination=True):
        conditions, parameters = build_ticket_filters(user)
        where = " WHERE " + " AND ".join(conditions) if conditions else ""
        db = get_db()
        total = db.execute(f"SELECT COUNT(*) AS count FROM tickets t{where}", parameters).fetchone()["count"]
        sql = f"""
            SELECT t.*, requester.name AS requester_name, requester.email AS requester_email,
                   assignee.name AS assignee_name
            FROM tickets t
            JOIN users requester ON requester.id = t.requester_id
            LEFT JOIN users assignee ON assignee.id = t.assignee_id
            {where}
            ORDER BY t.created_at DESC
        """
        if include_pagination:
            page = max(1, request.args.get("page", default=1, type=int))
            per_page = min(100, max(1, request.args.get("per_page", default=20, type=int)))
            rows = db.execute(sql + " LIMIT ? OFFSET ?", [*parameters, per_page, (page - 1) * per_page]).fetchall()
            return rows, total, page, per_page
        return db.execute(sql, parameters).fetchall(), total, None, None

    @app.get("/api/health")
    def health():
        return jsonify({"status": "ok", "service": "enterprise-ticket-api"})

    @app.post("/api/auth/login")
    def login():
        payload = request.get_json(silent=True) or {}
        email = str(payload.get("email", "")).strip().lower()
        password = str(payload.get("password", ""))
        user = get_db().execute("SELECT * FROM users WHERE lower(email) = ?", (email,)).fetchone()
        if user is None or not user["password_hash"] or not check_password_hash(user["password_hash"], password):
            return error("邮箱或密码错误", 401)
        safe_user = {key: user[key] for key in ("id", "name", "email", "department", "role")}
        token = serializer().dumps({"user_id": user["id"]})
        return jsonify({"token": token, "user": safe_user})

    @app.get("/api/auth/me")
    @auth_required()
    def auth_me():
        return jsonify(g.current_user)

    @app.get("/api/meta")
    @auth_required()
    def meta():
        user = g.current_user
        if user["role"] == "admin":
            users = get_db().execute("SELECT id, name, email, department, role FROM users ORDER BY role, name").fetchall()
        elif user["role"] == "agent":
            users = get_db().execute("SELECT id, name, email, department, role FROM users WHERE role IN ('agent', 'admin') ORDER BY role, name").fetchall()
        else:
            users = get_db().execute("SELECT id, name, email, department, role FROM users WHERE id = ?", (user["id"],)).fetchall()
        return jsonify({
            "statuses": STATUSES,
            "priorities": PRIORITIES,
            "departments": DEPARTMENTS,
            "categories": CATEGORIES,
            "users": [row_to_dict(row) for row in users],
            "current_user": user,
            "role_labels": ROLE_LABELS,
            "sla_hours": SLA_HOURS,
        })

    @app.get("/api/tickets")
    @auth_required()
    def list_tickets():
        rows, total, page, per_page = ticket_query(g.current_user)
        return jsonify({"items": [enrich_ticket(row) for row in rows], "page": page, "per_page": per_page, "total": total})

    @app.get("/api/tickets/<int:ticket_id>")
    @auth_required()
    def get_ticket(ticket_id):
        ticket = fetch_ticket(ticket_id, g.current_user)
        if ticket is None:
            return error("工单不存在或无权查看", 404)
        logs = get_db().execute(
            """
            SELECT l.*, u.name AS operator_name
            FROM ticket_logs l LEFT JOIN users u ON u.id = l.operator_id
            WHERE l.ticket_id = ? ORDER BY l.created_at DESC
            """,
            (ticket_id,),
        ).fetchall()
        enriched = enrich_ticket(ticket)
        return jsonify({
            "ticket": enriched,
            "logs": [row_to_dict(row) for row in logs],
            "permissions": {
                "allowed_transitions": allowed_transitions_for(enriched, g.current_user),
                "can_assign": g.current_user["role"] == "admin",
                "can_escalate": g.current_user["role"] == "admin" and enriched["sla_status"] == "OVERDUE",
            },
        })

    @app.post("/api/tickets")
    @auth_required("requester", "admin")
    def create_ticket():
        payload = request.get_json(silent=True) or {}
        required = ["title", "description", "department", "category", "priority"]
        missing = [field for field in required if payload.get(field) in (None, "")]
        if missing:
            return error("缺少必填字段", details=missing)
        if payload["priority"] not in PRIORITIES or payload["department"] not in DEPARTMENTS or payload["category"] not in CATEGORIES:
            return error("提交的数据包含无效选项")
        requester_id = g.current_user["id"] if g.current_user["role"] == "requester" else payload.get("requester_id")
        requester = get_db().execute("SELECT id, department FROM users WHERE id = ? AND role = 'requester'", (requester_id,)).fetchone()
        if requester is None:
            return error("提交人不存在")
        if g.current_user["role"] == "requester" and payload["department"] != g.current_user["department"]:
            return error("普通员工只能为本部门创建工单", 403)
        now = utc_now()
        due_at = now + timedelta(hours=SLA_HOURS[payload["priority"]])
        year_month = now.strftime("%Y%m")
        existing_numbers = get_db().execute("SELECT ticket_no FROM tickets WHERE ticket_no LIKE ?", (f"TKT-{year_month}-%",)).fetchall()
        next_number = max((int(row["ticket_no"].rsplit("-", 1)[-1]) for row in existing_numbers), default=0) + 1
        ticket_no = f"TKT-{year_month}-{next_number:04d}"
        try:
            cursor = get_db().execute(
                """
                INSERT INTO tickets (ticket_no, title, description, department, category, priority, status,
                    requester_id, assignee_id, created_at, updated_at, due_at)
                VALUES (?, ?, ?, ?, ?, ?, 'NEW', ?, NULL, ?, ?, ?)
                """,
                (ticket_no, payload["title"].strip(), payload["description"].strip(), payload["department"],
                 payload["category"], payload["priority"], requester_id, iso(now), iso(now), iso(due_at)),
            )
            ticket_id = cursor.lastrowid
            get_db().execute(
                "INSERT INTO ticket_logs (ticket_id, action, from_status, to_status, operator_id, note, created_at) VALUES (?, 'CREATED', NULL, 'NEW', ?, ?, ?)",
                (ticket_id, g.current_user["id"], payload.get("note", "工单已创建"), iso(now)),
            )
            get_db().commit()
        except sqlite3.IntegrityError as exc:
            return error("工单创建失败", details=[str(exc)])
        return jsonify(enrich_ticket(fetch_ticket(ticket_id, g.current_user))), 201

    @app.patch("/api/tickets/<int:ticket_id>")
    @auth_required("agent", "admin")
    def update_ticket(ticket_id):
        ticket = fetch_ticket(ticket_id, g.current_user)
        if ticket is None:
            return error("工单不存在或无权操作", 404)
        payload = request.get_json(silent=True) or {}
        editable = {"title", "description", "department", "category", "priority", "resolution"}
        if g.current_user["role"] == "admin":
            editable.add("assignee_id")
        updates = {key: value for key, value in payload.items() if key in editable}
        if not updates:
            return error("没有可更新的字段")
        updates["updated_at"] = iso(utc_now())
        assignments = ", ".join(f"{key} = ?" for key in updates)
        get_db().execute(f"UPDATE tickets SET {assignments} WHERE id = ?", [*updates.values(), ticket_id])
        get_db().execute(
            "INSERT INTO ticket_logs (ticket_id, action, operator_id, note, created_at) VALUES (?, 'UPDATED', ?, ?, ?)",
            (ticket_id, g.current_user["id"], payload.get("note", "更新工单信息"), updates["updated_at"]),
        )
        get_db().commit()
        return jsonify(enrich_ticket(fetch_ticket(ticket_id, g.current_user)))

    @app.post("/api/tickets/<int:ticket_id>/transition")
    @auth_required()
    def transition_ticket(ticket_id):
        ticket = fetch_ticket(ticket_id, g.current_user)
        if ticket is None:
            return error("工单不存在或无权操作", 404)
        payload = request.get_json(silent=True) or {}
        target = payload.get("status")
        if target not in allowed_transitions_for(ticket, g.current_user):
            return error("当前角色不能执行该状态变更", 403)
        assignee_id = ticket["assignee_id"]
        if g.current_user["role"] == "admin":
            assignee_id = payload.get("assignee_id", assignee_id)
        elif g.current_user["role"] == "agent" and target == "ASSIGNED":
            assignee_id = g.current_user["id"]
        if target in {"ASSIGNED", "IN_PROGRESS", "RESOLVED"} and not assignee_id:
            return error("进入处理流程前必须分配负责人")
        resolution = payload.get("resolution", ticket["resolution"])
        if target == "RESOLVED" and not str(resolution or "").strip():
            return error("解决工单时必须填写解决方案")
        now = utc_now()
        if target == "RESOLVED":
            resolved_at = iso(now)
        elif target == "CLOSED":
            resolved_at = ticket["resolved_at"] or iso(now)
        else:
            resolved_at = None
        get_db().execute(
            "UPDATE tickets SET status = ?, assignee_id = ?, resolution = ?, resolved_at = ?, updated_at = ? WHERE id = ?",
            (target, assignee_id, resolution, resolved_at, iso(now), ticket_id),
        )
        get_db().execute(
            "INSERT INTO ticket_logs (ticket_id, action, from_status, to_status, operator_id, note, created_at) VALUES (?, 'STATUS_CHANGED', ?, ?, ?, ?, ?)",
            (ticket_id, ticket["status"], target, g.current_user["id"], payload.get("note", "状态已更新"), iso(now)),
        )
        get_db().commit()
        return jsonify(enrich_ticket(fetch_ticket(ticket_id, g.current_user)))

    @app.post("/api/tickets/<int:ticket_id>/escalate")
    @auth_required("admin")
    def escalate_ticket(ticket_id):
        ticket = fetch_ticket(ticket_id, g.current_user)
        if ticket is None:
            return error("工单不存在", 404)
        enriched = enrich_ticket(ticket)
        if enriched["sla_status"] != "OVERDUE":
            return error("只有超时未完成工单可以升级")
        level = min(3, ticket["escalation_level"] + 1)
        now = iso(utc_now())
        get_db().execute("UPDATE tickets SET escalation_level = ?, updated_at = ? WHERE id = ?", (level, now, ticket_id))
        get_db().execute(
            "INSERT INTO ticket_logs (ticket_id, action, operator_id, note, created_at) VALUES (?, 'ESCALATED', ?, ?, ?)",
            (ticket_id, g.current_user["id"], f"SLA超时，工单升级至L{level}", now),
        )
        get_db().commit()
        return jsonify(enrich_ticket(fetch_ticket(ticket_id, g.current_user)))

    @app.get("/api/dashboard")
    @auth_required()
    def dashboard():
        db = get_db()
        now = utc_now()
        scope, parameters = user_scope(g.current_user)
        where = f" WHERE {scope}" if scope else ""
        prefix = " AND " if where else " WHERE "
        summary = row_to_dict(db.execute(
            f"""
            SELECT COUNT(*) AS total,
                SUM(CASE WHEN status IN ('RESOLVED', 'CLOSED') THEN 1 ELSE 0 END) AS resolved,
                SUM(CASE WHEN status NOT IN ('RESOLVED', 'CLOSED') AND due_at < ? THEN 1 ELSE 0 END) AS overdue,
                SUM(CASE WHEN escalation_level > 0 THEN 1 ELSE 0 END) AS escalated
            FROM tickets t{where}
            """,
            [iso(now), *parameters],
        ).fetchone())
        warning_conditions = []
        warning_parameters = []
        for priority, hours in SLA_HOURS.items():
            warning_conditions.append("(priority = ? AND due_at BETWEEN ? AND ?)")
            warning_parameters.extend([priority, iso(now), iso(now + timedelta(hours=hours * 0.25))])
        summary["warning"] = db.execute(
            f"SELECT COUNT(*) AS count FROM tickets t{where}{prefix}status NOT IN ('RESOLVED','CLOSED') AND ({' OR '.join(warning_conditions)})",
            [*parameters, *warning_parameters],
        ).fetchone()["count"]
        resolution_rows = db.execute(
            f"SELECT created_at, resolved_at, due_at FROM tickets t{where}{prefix}resolved_at IS NOT NULL", parameters
        ).fetchall()
        durations = [(parse_iso(row["resolved_at"]) - parse_iso(row["created_at"])).total_seconds() / 3600 for row in resolution_rows]
        met_count = sum(parse_iso(row["resolved_at"]) <= parse_iso(row["due_at"]) for row in resolution_rows)
        summary["resolution_rate"] = round(summary["resolved"] / summary["total"] * 100, 1) if summary["total"] else 0
        summary["avg_resolution_hours"] = round(sum(durations) / len(durations), 1) if durations else 0
        summary["sla_compliance_rate"] = round(met_count / len(resolution_rows) * 100, 1) if resolution_rows else 0

        def grouped(column):
            return [row_to_dict(row) for row in db.execute(
                f"SELECT {column} AS name, COUNT(*) AS value FROM tickets t{where} GROUP BY {column} ORDER BY value DESC", parameters
            ).fetchall()]

        since = iso(now - timedelta(days=29))
        created_rows = db.execute(
            f"SELECT substr(created_at,1,10) AS day, COUNT(*) AS value FROM tickets t{where}{prefix}created_at >= ? GROUP BY day",
            [*parameters, since],
        ).fetchall()
        resolved_rows = db.execute(
            f"SELECT substr(resolved_at,1,10) AS day, COUNT(*) AS value FROM tickets t{where}{prefix}resolved_at >= ? GROUP BY day",
            [*parameters, since],
        ).fetchall()
        created_by_day = {row["day"]: row["value"] for row in created_rows}
        resolved_by_day = {row["day"]: row["value"] for row in resolved_rows}
        trend = []
        for offset in range(29, -1, -1):
            day = (now - timedelta(days=offset)).strftime("%Y-%m-%d")
            trend.append({"day": day, "created": created_by_day.get(day, 0), "resolved": resolved_by_day.get(day, 0)})
        return jsonify({
            "summary": summary,
            "by_status": grouped("status"),
            "by_priority": grouped("priority"),
            "by_category": grouped("category"),
            "by_department": grouped("department"),
            "trend": trend,
            "generated_at": iso(now),
        })

    @app.get("/api/reports/tickets.<file_format>")
    @auth_required()
    def export_tickets(file_format):
        if file_format not in {"csv", "xlsx"}:
            return error("仅支持CSV和Excel格式", 404)
        rows, _, _, _ = ticket_query(g.current_user, include_pagination=False)
        tickets = [enrich_ticket(row) for row in rows]
        headers = ["工单编号", "问题标题", "提交部门", "问题类型", "优先级", "状态", "提交人", "负责人", "SLA状态", "升级级别", "创建时间", "截止时间", "解决时间"]
        data = [[
            item["ticket_no"], item["title"], item["department"], item["category"], PRIORITY_LABELS[item["priority"]],
            STATUS_LABELS[item["status"]], item["requester_name"], item["assignee_name"] or "未分配", item["sla_status"],
            item["escalation_level"], item["created_at"], item["due_at"], item["resolved_at"] or "",
        ] for item in tickets]
        if file_format == "csv":
            stream = StringIO()
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(data)
            output = BytesIO(stream.getvalue().encode("utf-8-sig"))
            return send_file(output, mimetype="text/csv", as_attachment=True, download_name="tickets-report.csv")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "工单明细"
        sheet.append(headers)
        for row in data:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        fill = PatternFill("solid", fgColor="0F5BA7")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
        widths = [20, 34, 14, 14, 10, 12, 14, 14, 14, 10, 24, 24, 24]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="tickets-report.xlsx")

    if frontend_dist.exists():
        @app.get("/")
        def frontend_index():
            return send_from_directory(frontend_dist, "index.html")

        @app.get("/<path:asset_path>")
        def frontend_assets(asset_path):
            target = frontend_dist / asset_path
            if target.is_file():
                return send_from_directory(frontend_dist, asset_path)
            return send_from_directory(frontend_dist, "index.html")

    return app


if __name__ == "__main__":
    application = create_app()
    application.run(host="0.0.0.0", port=5000, debug=True)
